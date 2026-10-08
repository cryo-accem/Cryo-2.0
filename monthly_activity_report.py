import calendar
import datetime
from collections.abc import Iterable, Mapping
from decimal import Decimal
from io import BytesIO

from docx import Document as create_document
from docx.document import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table
from docx.shared import Inches, Pt, RGBColor


ZERO = Decimal("0")


def _date_value(value: object) -> datetime.date | None:
    if not value:
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    try:
        return datetime.date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _decimal(value: object) -> Decimal:
    try:
        return Decimal(str(value or 0))
    except (ValueError, ArithmeticError):
        return ZERO


def _grids(row: Mapping[str, object]) -> int:
    return int(_decimal(row.get("number_of_grids") or row.get("actual_grids") or row.get("grids")))


def _grid_source(row: Mapping[str, object]) -> str:
    source = str(row.get("grid_source") or "").strip().casefold().replace("-", "_").replace(" ", "_")
    return "facility" if source in {"facility", "facility_provided"} else "outside"


def _money(value: Decimal) -> str:
    return f"INR {value:,.2f}"


def _format_date(value: object) -> str:
    date_value = _date_value(value)
    return date_value.strftime("%d %b %Y") if date_value else "-"


def _row_revenue(row: Mapping[str, object]) -> Decimal:
    return _decimal(row.get("total_billed") or row.get("grand_total"))


def _add_table(
    document: Document,
    headers: tuple[str, ...],
    rows: Iterable[tuple[object, ...]],
    widths: tuple[float, ...] | None = None,
) -> Table:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            cell.text = str(value)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    return table


def _add_service_table(
    document: Document, title: str, events: list[dict[str, object]], service: str
) -> None:
    document.add_heading(title, level=2)
    if not events:
        document.add_paragraph("No completed records in this period.")
        return
    rows = []
    for row in events:
        clipped = int(_decimal(row.get("clipped_grids")))
        rows.append((
            _format_date(row.get("_report_date")),
            row.get("user_name") or "-",
            row.get("pi_name") or "-",
            row.get("sample_name") or "-",
            _grids(row),
            _grid_source(row),
            clipped if service != "Freezing" else "—",
            _money(_row_revenue(row)),
        ))
    _add_table(
        document,
        ("Date", "User", "PI", "Sample", "Grids", "Grid source", "Clipped", "Billed incl. GST"),
        rows,
    )


def build_monthly_activity_report(
    start: datetime.date,
    end: datetime.date,
    data_collection: Iterable[Mapping[str, object]],
    screening: Iterable[Mapping[str, object]],
    freezing: Iterable[Mapping[str, object]],
    historical_revenue: Iterable[tuple[str, Decimal]],
) -> BytesIO:
    """Build an editable Word report for completed service activity in a date range."""
    events_by_service = {
        "Data collection": [],
        "Screening": [],
        "Freezing": [],
    }
    for service, source_rows in (
        ("Data collection", data_collection),
        ("Screening", screening),
        ("Freezing", freezing),
    ):
        date_key = "completed_at" if service == "Freezing" else "completion_date"
        for source_row in source_rows:
            row = dict(source_row)
            event_date = _date_value(row.get(date_key))
            if event_date and start <= event_date <= end:
                row["_report_date"] = event_date
                events_by_service[service].append(row)
        events_by_service[service].sort(key=lambda row: (row["_report_date"], str(row.get("id") or "")))

    clipped_events = [
        (service, row)
        for service, events in events_by_service.items()
        for row in events
        if int(_decimal(row.get("clipped_grids"))) > 0
    ]
    all_events = [row for events in events_by_service.values() for row in events]
    facility_grids = sum(_grids(row) for row in all_events if _grid_source(row) == "facility")
    outside_grids = sum(_grids(row) for row in all_events if _grid_source(row) == "outside")
    clipped_grids = sum(int(_decimal(row.get("clipped_grids"))) for _, row in clipped_events)
    current_billed = sum((_row_revenue(row) for row in all_events), ZERO)
    received = sum((_decimal(row.get("amount_received")) for row in all_events), ZERO)
    historical_by_month = {
        month: _decimal(amount)
        for month, amount in historical_revenue
        if start.strftime("%Y-%m") <= month <= end.strftime("%Y-%m")
    }
    historical_total = sum(historical_by_month.values(), ZERO)
    gross_revenue = current_billed + historical_total

    document = create_document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(9)

    title = document.add_heading("ACCEM Monthly Activity & Revenue Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(
        f"Reporting period: {start:%d %B %Y} to {end:%d %B %Y}"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("Period overview", level=1)
    _add_table(document, ("Measure", "Total"), (
        ("Completed freezing records", len(events_by_service["Freezing"])),
        ("Completed clipping records", len(clipped_events)),
        ("Completed screening records", len(events_by_service["Screening"])),
        ("Completed data collection records", len(events_by_service["Data collection"])),
        ("Clipped grids", clipped_grids),
        ("Facility-provided grids used", facility_grids),
        ("Outside/user-provided grids used", outside_grids),
        ("Completed-record billed revenue (incl. GST)", _money(current_billed)),
        ("Legacy historical revenue (pre-tax summary)", _money(historical_total)),
        ("Total revenue represented", _money(gross_revenue)),
        ("Payments received on completed records", _money(received)),
        ("Outstanding on completed records", _money(max(current_billed - received, ZERO))),
    ))

    document.add_heading("Month-by-month summary", level=1)
    month_rows = []
    month_cursor = start.replace(day=1)
    while month_cursor <= end:
        month_end = month_cursor.replace(day=calendar.monthrange(month_cursor.year, month_cursor.month)[1])
        month_start = max(start, month_cursor)
        effective_end = min(end, month_end)
        in_month = [
            row for row in all_events
            if month_start <= row["_report_date"] <= effective_end
        ]
        month_history = historical_by_month.get(month_cursor.strftime("%Y-%m"), ZERO)
        month_rows.append((
            month_cursor.strftime("%B %Y"),
            sum(1 for row in events_by_service["Freezing"] if month_start <= row["_report_date"] <= effective_end),
            sum(1 for _, row in clipped_events if month_start <= row["_report_date"] <= effective_end),
            sum(1 for row in events_by_service["Screening"] if month_start <= row["_report_date"] <= effective_end),
            sum(1 for row in events_by_service["Data collection"] if month_start <= row["_report_date"] <= effective_end),
            sum(_grids(row) for row in in_month if _grid_source(row) == "facility"),
            sum(_grids(row) for row in in_month if _grid_source(row) == "outside"),
            _money(sum((_row_revenue(row) for row in in_month), ZERO) + month_history),
        ))
        month_cursor = (month_cursor.replace(day=28) + datetime.timedelta(days=4)).replace(day=1)
    _add_table(
        document,
        ("Month", "Freezing", "Clipping", "Screening", "Data collection", "Facility grids", "Outside grids", "Revenue"),
        month_rows,
    )

    document.add_heading("Completed activity records", level=1)
    _add_service_table(document, "Freezing", events_by_service["Freezing"], "Freezing")
    _add_service_table(document, "Screening", events_by_service["Screening"], "Screening")
    _add_service_table(document, "Data collection", events_by_service["Data collection"], "Data collection")

    document.add_heading("Clipping records", level=2)
    if clipped_events:
        clipping_rows = [
            (
                service,
                _format_date(row.get("_report_date")),
                row.get("user_name") or "-",
                row.get("pi_name") or "-",
                row.get("sample_name") or "-",
                int(_decimal(row.get("clipped_grids"))),
                _money(_decimal(row.get("clipping_charge"))),
            )
            for service, row in clipped_events
        ]
        _add_table(
            document,
            ("Service", "Date", "User", "PI", "Sample", "Grids clipped", "Clipping charge"),
            clipping_rows,
        )
    else:
        document.add_paragraph("No completed records with clipped grids in this period.")

    if historical_total:
        document.add_heading("Historical revenue note", level=1)
        document.add_paragraph(
            "Historical revenue is included as a pre-tax monthly aggregate from legacy records. "
            "Those records do not contain itemized service, grid-usage, GST, or payment details "
            "and are therefore not included in the activity lists above."
        )
    document.add_paragraph(
        "Grid counts represent completed service uses recorded by the facility; the same grids "
        "may appear in more than one service event. Outside grids include user-provided or "
        "unclassified legacy grid sources."
    )

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
